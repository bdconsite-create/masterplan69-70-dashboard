#!/usr/bin/env python3
"""Build the dashboard's static data and photo assets from the source workbook."""

from __future__ import annotations

import json
import re
import shutil
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from posixpath import dirname, join, normpath
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook


DAY_ZERO = datetime(1899, 12, 30)


def iso(value):
    if value in (None, ""):
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        return (DAY_ZERO + timedelta(days=float(value))).strftime("%Y-%m-%d")
    text = str(value).strip()
    match = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", text)
    if match:
        day, month, year = map(int, match.groups())
        if year > 2400:
            year -= 543
        return f"{year:04d}-{month:02d}-{day:02d}"
    return text[:10] if re.match(r"\d{4}-\d{2}-\d{2}", text) else ""


def clean(value):
    return "" if value is None else str(value).strip()


def sheet_xml_path(archive: ZipFile, sheet_name: str):
    workbook = etree.fromstring(archive.read("xl/workbook.xml"))
    rels = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_targets = {
        node.get("Id"): node.get("Target")
        for node in rels.xpath('//*[local-name()="Relationship"]')
    }
    for node in workbook.xpath('//*[local-name()="sheet"]'):
        if node.get("name") != sheet_name:
            continue
        rel_id = node.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        target = rel_targets.get(rel_id, "")
        if target.startswith("/"):
            return target.lstrip("/")
        return normpath(join(dirname("xl/workbook.xml"), target))
    raise KeyError(f"sheet not found: {sheet_name}")


def header_columns(sheet, row=4):
    return {
        re.sub(r"[\s/]+", "", clean(sheet.cell(row, col).value)): col
        for col in range(1, sheet.max_column + 1)
        if clean(sheet.cell(row, col).value)
    }


def find_header(headers, predicate, fallback=None):
    for label, column in headers.items():
        if predicate(label):
            return column
    if fallback is not None:
        return fallback
    raise KeyError("required header not found")


def extract_rich_photos(
    workbook_path: Path,
    output_dir: Path,
    sheet_name: str,
    photo_rows: dict[int, dict],
    photo_columns: set[int],
    filename_prefix: str = "",
):
    photos = []
    with ZipFile(workbook_path) as archive:
        sheet = etree.fromstring(archive.read(sheet_xml_path(archive, sheet_name)))
        values = etree.fromstring(archive.read("xl/richData/rdrichvalue.xml"))
        links = etree.fromstring(archive.read("xl/richData/richValueRel.xml"))
        link_rels = etree.fromstring(archive.read("xl/richData/_rels/richValueRel.xml.rels"))

        rich_indexes = [
            int(node.xpath('./*[local-name()="v"][1]/text()')[0])
            for node in values.xpath('//*[local-name()="rv"]')
        ]
        rel_ids = [
            node.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            for node in links.xpath('//*[local-name()="rel"]')
        ]
        media_by_id = {
            node.get("Id"): normpath(join(dirname("xl/richData/richValueRel.xml"), node.get("Target")))
            for node in link_rels.xpath('//*[local-name()="Relationship"]')
        }

        for cell in sheet.xpath('//*[local-name()="c"][@vm]'):
            ref = cell.get("r")
            match = re.fullmatch(r"([A-Z]+)(\d+)", ref)
            if not match:
                continue
            letters, row_text = match.groups()
            row = int(row_text)
            col = 0
            for char in letters:
                col = col * 26 + ord(char) - 64
            if row < 5 or col not in photo_columns or row not in photo_rows:
                continue

            value_index = int(cell.get("vm")) - 1
            if not 0 <= value_index < len(rich_indexes):
                continue
            rel_index = rich_indexes[value_index]
            if not 0 <= rel_index < len(rel_ids):
                continue
            media_path = media_by_id.get(rel_ids[rel_index])
            if not media_path or media_path not in archive.namelist():
                continue

            row_data = photo_rows[row]
            slot = sorted(photo_columns).index(col) + 1
            suffix = Path(media_path).suffix.lower() or ".jpg"
            file_task = row_data["taskId"] if re.fullmatch(r"T\d+", row_data["taskId"]) else "UNASSIGNED"
            filename = f"{filename_prefix}{file_task}_{row_data['date']}_photo-{slot}{suffix}"
            if (output_dir / filename).exists():
                filename = f"{filename_prefix}{file_task}_{row_data['date']}_r{row}_photo-{slot}{suffix}"
            (output_dir / filename).write_bytes(archive.read(media_path))
            photos.append({
                "id": f"base-{row}-{col}",
                "vessel": row_data["vessel"],
                "taskId": row_data["taskId"],
                "date": row_data["date"],
                "detail": row_data["detail"],
                "slot": slot,
                "url": f"photos/{filename}",
            })
    return photos


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: update_from_excel.py INPUT.xlsx DIST_DIR")
    workbook_path = Path(sys.argv[1]).resolve()
    dist_dir = Path(sys.argv[2]).resolve()
    photo_dir = dist_dir / "photos"
    if photo_dir.exists():
        shutil.rmtree(photo_dir)
    photo_dir.mkdir(parents=True)

    # Normal mode makes repeated cell lookups O(1).  Read-only mode rescans the
    # wide Gantt sheet for every lookup and is dramatically slower here.
    book = load_workbook(workbook_path, data_only=True, read_only=False)
    master = book["mater69-70"]
    master_headers = header_columns(master)
    order_col = find_header(master_headers, lambda x: x == "ลำดับ", 1)
    task_col = find_header(master_headers, lambda x: x.lower() == "taskid", 2)
    vessel_col = find_header(master_headers, lambda x: x.startswith("เรือ"), 3)
    section_col = find_header(master_headers, lambda x: x.startswith("แผนก"), 4)
    status_col = find_header(master_headers, lambda x: x == "สถานะ", 5)
    work_col = find_header(master_headers, lambda x: x == "รายละเอียดงาน", 6)
    start_col = find_header(master_headers, lambda x: x == "เริ่ม", 9)
    end_col = find_header(master_headers, lambda x: x == "เสร็จ", 10)
    kind_col = find_header(master_headers, lambda x: "แผนงาน" in x and "ผลงาน" in x, 11)
    tasks, summaries = [], []
    for row in range(7, master.max_row + 1):
        if clean(master.cell(row, kind_col).value) != "Plan":
            continue
        actual_row = row + 1
        task_id = clean(master.cell(row, task_col).value)
        vessel = clean(master.cell(row, vessel_col).value)
        work = clean(master.cell(row, work_col).value)
        if work.startswith("แผนงานรวม"):
            summaries.append({
                "vessel": vessel,
                "planStart": iso(master.cell(row, start_col).value),
                "planEnd": iso(master.cell(row, end_col).value),
                "actualStart": iso(master.cell(actual_row, start_col).value),
                "actualEnd": iso(master.cell(actual_row, end_col).value),
            })
        elif re.fullmatch(r"T\d+", task_id) and work:
            tasks.append({
                "id": task_id,
                "order": clean(master.cell(row, order_col).value),
                "vessel": vessel,
                "section": clean(master.cell(row, section_col).value),
                "status": clean(master.cell(row, status_col).value),
                "work": work,
                "planStart": iso(master.cell(row, start_col).value),
                "planEnd": iso(master.cell(row, end_col).value),
                "actualStart": iso(master.cell(actual_row, start_col).value),
                "actualEnd": iso(master.cell(actual_row, end_col).value),
            })

    daily = book["Daily Photos"]
    daily_headers = header_columns(daily)
    daily_vessel_col = find_header(daily_headers, lambda x: x == "เรือ", 1)
    daily_task_col = find_header(daily_headers, lambda x: x.lower() == "taskid", 2)
    daily_date_col = find_header(daily_headers, lambda x: x == "วันที่", 3)
    daily_detail_col = find_header(daily_headers, lambda x: x.startswith("รายละเอียด"), 4)
    photo_columns = {
        column for label, column in daily_headers.items()
        if label.startswith("รูปที่") or label.startswith("รูปภาพ")
    }
    if not photo_columns:
        photo_columns = set(range(5, 10))
    photo_rows = {}
    for row in range(5, daily.max_row + 1):
        task_id = clean(daily.cell(row, daily_task_col).value)
        vessel = clean(daily.cell(row, daily_vessel_col).value)
        photo_date = iso(daily.cell(row, daily_date_col).value)
        detail = clean(daily.cell(row, daily_detail_col).value)
        if not any((task_id, vessel, photo_date, detail)):
            continue
        photo_rows[row] = {
            "vessel": vessel,
            "taskId": task_id or "ยังไม่ระบุ Task ID",
            "date": photo_date,
            "detail": detail,
        }

    photos = extract_rich_photos(workbook_path, photo_dir, "Daily Photos", photo_rows, photo_columns)
    photos.sort(key=lambda p: (p["date"], p["vessel"], p["taskId"], p["slot"]), reverse=True)

    detail_items, detail_photos = [], []
    if "รูปภาพรายละเอียดงาน" in book.sheetnames:
        detail_sheet = book["รูปภาพรายละเอียดงาน"]
        detail_headers = header_columns(detail_sheet, row=1)
        detail_vessel_col = find_header(detail_headers, lambda x: x == "เรือ", 1)
        detail_task_col = find_header(detail_headers, lambda x: x.lower() == "taskid", 2)
        detail_date_col = find_header(detail_headers, lambda x: x == "วันที่", 3)
        detail_text_col = find_header(detail_headers, lambda x: x.startswith("รายละเอียด"), 4)
        detail_photo_columns = {
            column for label, column in detail_headers.items()
            if label.startswith("รูปภาพ") or label.startswith("รูปที่")
        }
        if not detail_photo_columns:
            detail_photo_columns = set(range(5, 15))
        detail_rows = {}
        for row in range(2, detail_sheet.max_row + 1):
            task_id = clean(detail_sheet.cell(row, detail_task_col).value)
            vessel = clean(detail_sheet.cell(row, detail_vessel_col).value)
            detail_date = iso(detail_sheet.cell(row, detail_date_col).value)
            detail_text = clean(detail_sheet.cell(row, detail_text_col).value)
            if not any((task_id, vessel, detail_date, detail_text)):
                continue
            row_data = {
                "vessel": vessel,
                "taskId": task_id or "ยังไม่ระบุ Task ID",
                "date": detail_date,
                "detail": detail_text,
            }
            detail_rows[row] = row_data
            if task_id and vessel and detail_text:
                detail_items.append(row_data)
        detail_photos = extract_rich_photos(
            workbook_path,
            photo_dir,
            "รูปภาพรายละเอียดงาน",
            detail_rows,
            detail_photo_columns,
            filename_prefix="DETAIL_",
        )
        detail_photos.sort(key=lambda p: (p["vessel"], p["taskId"], p["slot"]))

    output = {
        "tasks": tasks,
        "summaries": summaries,
        "photos": photos,
        "detailItems": detail_items,
        "detailPhotos": detail_photos,
    }
    (dist_dir / "data.json").write_text(json.dumps(output, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({
        "tasks": len(tasks),
        "summaries": len(summaries),
        "photos": len(photos),
        "detailItems": len(detail_items),
        "detailPhotos": len(detail_photos),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
